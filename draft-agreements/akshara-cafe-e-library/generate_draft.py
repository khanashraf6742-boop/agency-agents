#!/usr/bin/env python3
"""
Generate: Draft License Agreement for Approval
Akshara Café at E-Library, NITK Surathkal
Basis: Commercial Establishment Committee (CEC) Minutes dated 23-09-2026,
Ref. No. 07/NITK/AC-1/2026-27/COM-ESTB/A6 dated 24-09-2026, Item No. 4.

Outputs a professionally formatted .docx (editable, for approval routing).
Placeholders are shown as [[ ... ]] and highlighted in yellow in the DOCX.
"""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from docx.dml.color import ColorFormat

OUT = "/home/user/agency-agents/draft-agreements/akshara-cafe-e-library/01_Draft_License_Agreement_Akshara_Cafe_E-Library_For_Approval.docx"

NAVY = RGBColor(0x1F, 0x2E, 0x6E)
MAROON = RGBColor(0x7A, 0x1F, 0x1F)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x55, 0x55)
HEADER_BG = "1F2E6E"
LIGHT_BG = "EEF1F8"
LIGHT_GREY_BG = "F2F2F2"

doc = Document()

# ---------- Page setup ----------
for section in doc.sections:
    section.page_width = Inches(8.27)   # A4
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)
style.font.color.rgb = DARK
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.line_spacing = 1.15

for i, (sz, col) in enumerate([(16, NAVY), (13, NAVY), (11.5, DARK)], start=1):
    hs = doc.styles[f'Heading {i}']
    hs.font.name = 'Times New Roman'
    hs.font.size = Pt(sz)
    hs.font.color.rgb = col
    hs.font.bold = True
    hs.paragraph_format.space_before = Pt(10 if i == 1 else 8)
    hs.paragraph_format.space_after = Pt(4)
    hs.paragraph_format.keep_with_next = True

# ---------- Header / Footer ----------
def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instr = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fld2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)

sec = doc.sections[0]
hdr = sec.header
hp = hdr.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = hp.add_run("DRAFT  FOR  APPROVAL  —  NOT  FOR  EXECUTION")
r.font.name = 'Times New Roman'; r.font.size = Pt(8); r.font.bold = True
r.font.color.rgb = MAROON
hp.paragraph_format.space_after = Pt(2)

ftr = sec.footer
fp = ftr.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Draft License Agreement — Akshara Café, E-Library, NITK Surathkal  |  Basis: CEC Minutes dated 23-09-2026, Item No. 4  |  Page ")
r.font.name = 'Times New Roman'; r.font.size = Pt(7.5); r.font.color.rgb = GREY
add_page_number(fp)
for run in fp.runs:
    run.font.name = 'Times New Roman'; run.font.size = Pt(7.5); run.font.color.rgb = GREY

# ---------- Helpers ----------
def parse_runs(paragraph, text, base_size=11, base_bold=False, base_italic=False,
               align=None, space_after=4, space_before=0, left_indent=None, first_line=None):
    """Supports **bold**, *italic*, [[placeholder - yellow highlight]]"""
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(space_before)
    if left_indent is not None:
        paragraph.paragraph_format.left_indent = Inches(left_indent)
    if first_line is not None:
        paragraph.paragraph_format.first_line_indent = Inches(first_line)
    # tokenise
    parts = re.split(r'(\[\[.*?\]\]|\*\*.*?\*\*|\*[^*]+?\*)', text)
    for part in parts:
        bold = base_bold; italic = base_italic; hl = False; t = part
        if part.startswith('[[') and part.endswith(']]'):
            t = part[2:-2]; hl = True; bold = True
        elif part.startswith('**') and part.endswith('**'):
            t = part[2:-2]; bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            t = part[1:-1]; italic = True
        run = paragraph.add_run(t)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(base_size)
        run.bold = bold; run.italic = italic
        run.font.color.rgb = DARK
        if hl:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return paragraph

def p(text, **kw):
    para = doc.add_paragraph()
    defaults = dict(align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    defaults.update(kw)
    return parse_runs(para, text, **defaults)

def p_center(text, size=11, bold=False, color=None, space_after=4, space_before=0, italic=False):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(space_before)
    run = para.add_run(text)
    run.font.name = 'Times New Roman'; run.font.size = Pt(size); run.bold = bold; run.italic = italic
    if color: run.font.color.rgb = color
    return para

def h1(text, center=False):
    para = doc.add_heading(level=1)
    if center: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text); run.font.name = 'Times New Roman'
    return para

def h2(text):
    para = doc.add_heading(level=2)
    run = para.add_run(text); run.font.name = 'Times New Roman'
    return para

def h3(text):
    para = doc.add_heading(level=3)
    run = para.add_run(text); run.font.name = 'Times New Roman'
    return para

def bullet(text, level=0, bullet_char="•"):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.28 + level * 0.22)
    para.paragraph_format.first_line_indent = Inches(-0.18)
    parse_runs(para, f"{bullet_char}  {text}", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2)
    return para

def numbered(label, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.32)
    para.paragraph_format.first_line_indent = Inches(-0.32)
    parse_runs(para, f"**{label}**  {text}", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=3)
    return para

def subclause(label, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.55)
    para.paragraph_format.first_line_indent = Inches(-0.3)
    parse_runs(para, f"{label}  {text}", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2)
    return para

def line():
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(8)
    pPr = para._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="12" w:space="1" w:color="1F2E6E"/>'
        f'</w:pBdr>')
    pPr.append(pBdr)
    return para

def shade(cell, hexcolor):
    tblCell = cell._tc
    tblCellProperties = tblCell.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hexcolor}"/>')
    tblCellProperties.append(shd)

def set_cell(cell, text, bold=False, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, color=None, white=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if cell.paragraphs:
        para = cell.paragraphs[0]
    else:
        para = cell.add_paragraph()
    para.alignment = align
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.space_before = Pt(2)
    # support [[ ]] highlight and ** **
    parts = re.split(r'(\[\[.*?\]\]|\*\*.*?\*\*)', text)
    for part in parts:
        b = bold; hl = False; t = part
        if part.startswith('[[') and part.endswith(']]'):
            t = part[2:-2]; hl = True; b = True
        elif part.startswith('**') and part.endswith('**'):
            t = part[2:-2]; b = True
        run = para.add_run(t)
        run.font.name = 'Times New Roman'; run.font.size = Pt(size); run.bold = b
        if white: run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        elif color: run.font.color.rgb = color
        else: run.font.color.rgb = DARK
        if hl: run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return cell

def table(rows, widths=None, header=True, size=10, col_aligns=None):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        is_hdr = header and ri == 0
        for ci, val in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT
            if col_aligns and ci < len(col_aligns):
                align = col_aligns[ci]
            if is_hdr:
                shade(cells[ci], HEADER_BG)
                set_cell(cells[ci], val, bold=True, size=size, align=WD_ALIGN_PARAGRAPH.CENTER, white=True)
            else:
                if ri % 2 == 0:
                    shade(cells[ci], LIGHT_BG)
                set_cell(cells[ci], val, size=size, align=align)
    if widths:
        for ci, w in enumerate(widths):
            for ri in range(len(rows)):
                t.cell(ri, ci).width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def sig_table(rows):
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for left, right in rows:
        cells = t.add_row().cells
        for ci, block in enumerate([left, right]):
            cell = cells[ci]
            cell.width = Inches(3.28)
            cell.paragraphs[0].clear()
            for j, ln in enumerate(block):
                para = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                para.paragraph_format.space_after = Pt(2)
                parse_runs(para, ln, base_size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def letterhead():
    p_center("लेखा अनुभाग–I  /  Accounts Section–I", size=10.5, bold=True, color=NAVY, space_after=1)
    p_center("राष्ट्रीय प्रौद्योगिकी संस्थान कर्नाटक, सुरत्कल, मंगलूरु – ५७५०२५", size=13, bold=True, color=NAVY, space_after=1)
    p_center("(राष्ट्रीय महत्व का संस्थान, एनआईटीएसईआर अधिनियम के तहत स्थापित, भारत सरकार)", size=9.5, color=GREY, space_after=1)
    p_center("National Institute of Technology Karnataka, Surathkal, Mangaluru – 575 025", size=13, bold=True, color=NAVY, space_after=1)
    p_center("(An Institute of National Importance, Established under NITSER Act, Government of India)", size=9.5, color=GREY, space_after=1)
    p_center("दूरभाष / Phone: +91-824-2473020 (O)   |   वेबसाइट / Website: http://www.nitk.ac.in", size=9.5, color=GREY, space_after=2)
    line()

def pagebreak():
    doc.add_page_break()

# ============================================================
# COVER / LETTERHEAD
# ============================================================
letterhead()

# Ref / Date row as table
t = doc.add_table(rows=1, cols=2)
t.autofit = False
c0 = t.cell(0, 0); c1 = t.cell(0, 1)
c0.width = Inches(4.6); c1.width = Inches(1.97)
set_cell(c0, "Ref. No.: [[07/NITK/AC-1/2026-27/COM-ESTB/____ ]]", size=10.5)
set_cell(c1, "Date: [[____-____-2026]]", size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
for row in t.rows:
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/><w:right w:val="nil"/><w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tcBorders>')
        tcPr.append(borders)
doc.add_paragraph().paragraph_format.space_after = Pt(4)

# Draft stamp box
t2 = doc.add_table(rows=1, cols=1)
t2.style = 'Table Grid'; t2.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = t2.cell(0, 0)
shade(cell, "FFF2CC")
cell.paragraphs[0].clear()
para = cell.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = para.add_run("DRAFT FOR APPROVAL  •  अनुमोदन हेतु प्रारूप")
r.font.name = 'Times New Roman'; r.font.size = Pt(13); r.bold = True; r.font.color.rgb = MAROON
para2 = cell.add_paragraph(); para2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = para2.add_run("This draft has no legal effect until approved by the Competent Authority and executed on requisite stamp paper. "
                  "Highlighted portions are to be filled / confirmed before execution.")
r.font.name = 'Times New Roman'; r.font.size = Pt(8.5); r.italic = True; r.font.color.rgb = DARK
doc.add_paragraph().paragraph_format.space_after = Pt(4)

p_center("DRAFT LICENSE AGREEMENT", size=18, bold=True, color=NAVY, space_after=1, space_before=4)
p_center("for Operation of Coffee Machine at “Akshara Café”, E-Library Premises", size=12.5, bold=True, color=DARK, space_after=2)
p_center("Between", size=11, italic=True, space_after=1)
p_center("National Institute of Technology Karnataka, Surathkal  (Licensor)", size=11.5, bold=True, space_after=1)
p_center("and", size=11, italic=True, space_after=1)
p_center("NITK Consumer Co-operative Society Ltd.  (Licensee)", size=11.5, bold=True, space_after=8)

table([
    ["Particulars", "Details (as resolved by CEC / to be confirmed)"],
    ["Basis", "Minutes of the Commercial Establishment Committee (CEC) Meeting held on **23-09-2026**, Ref. No. **07/NITK/AC-1/2026-27/COM-ESTB/A6 dated 24-09-2026 — Item No. 4**"],
    ["Subject", "Licence to the NITK Consumer Co-operative Society Ltd. to install and operate a coffee machine at **“Akshara Café” in the E-Library premises**, NITK Surathkal"],
    ["Licensee", "NITK Consumer Co-operative Society Ltd., [[Regn. No. & Registered Office address to be inserted]]"],
    ["License Fee (Rent)", "**Rs. 3,000/- (Rupees Three Thousand only) per month** [[+ GST / other taxes, if and as applicable — for confirmation]]"],
    ["Security Deposit", "**Rs. 30,000/- (Rupees Thirty Thousand only)**, interest-free and refundable (subject to adjustments)"],
    ["Term", "**One (1) year** from the date of handing over the premises / commencement, as recorded in Schedule-I"],
    ["Water Charges", "At **standard rates until a water meter is installed**; thereafter on actual metered consumption at Institute rates"],
    ["Premises", "Earmarked space for “Akshara Café” in the E-Library, measuring [[____ sq. ft.]] — details and plan at **Schedule-I** (to be jointly measured and signed)"],
    ["Prepared by", "Accounts Section-I — for approval of the Competent Authority and subsequent forwarding to the Society, in terms of Resolution 4(b)"],
], widths=[1.75, 4.82], size=10)

p("**Contents of this DFA file:**  **Part-A** — Note for Approval (background, salient terms, decision sought and approval routing);  "
  "**Part-B** — Draft License Agreement (execution text, with Schedules I–III and Annexure-A);  "
  "**Part-C** — Draft forwarding letter to the NITK Consumer Co-operative Society Ltd. for issue after approval, in terms of Resolution 4(b).", base_size=10)

p("*Note: Clauses which are not expressly covered in the CEC Minutes (e.g., mode/ due date of payment, electricity charges, food-safety & hygiene compliance, "
  "Library discipline, termination notice, penal charges for overstay, dispute resolution) have been proposed on the basis of standard Institute practice for "
  "commercial establishments and are specifically flagged in Part-A, Para 6 for the approval / modification of the Competent Authority. "
  "The draft may be got vetted by the Institute’s Legal Advisor / Standing Counsel before execution, if so desired.*", base_size=9.5)

# ============================================================
# PART A — NOTE FOR APPROVAL
# ============================================================
pagebreak()
p_center("PART – A", size=12, bold=True, color=MAROON, space_after=1)
h1("Note for Approval", center=True)
line()

h2("1.  Reference")
numbered("1.1", "Minutes of the Meeting of the Commercial Establishment Committee (CEC) held on **23-09-2026 (Wednesday) at 3:00 P.M.** in the CCMT Meeting Hall, NITK Surathkal, under the Chairmanship of the Dean (FSW), issued vide Ref. No. **07/NITK/AC-1/2026-27/COM-ESTB/A6 dated 24-09-2026** — **Item No. 4: Proposal for Starting “Akshara Café” at E-Library** (extract reproduced at **Annexure-A** to Part-B).")
numbered("1.2", "Request of the NITK Consumer Co-operative Society Ltd. to run a coffee machine at “Akshara Café” in the Library premises, as considered by the CEC. [[Insert Society’s letter No. & date, if available]]")
numbered("1.3", "Copy of the Minutes was, inter alia, endorsed to the Librarian, NITK Surathkal, and to the NITK Consumer Co-operative Society Ltd., for information and necessary action.")

h2("2.  Resolution under Item No. 4")
p("The Committee resolved, in substance, as under:")
bullet("**Resolution 4(a):** The request of the NITK Consumer Co-operative Society Ltd. to run a coffee machine at “Akshara Café” in the Library premises is **approved** on the following terms: (i) Rent — **Rs. 3,000/- per month**; (ii) Security Deposit — **Rs. 30,000/-**; (iii) Contract Period — **One (1) year**; (iv) Water Charges — **at standard rates until a water meter is installed**.")
bullet("**Resolution 4(b):** The Consumer Co-operative Society **shall sign an agreement** with the Institute. The **draft agreement copy shall be sent to the Society**.")

h2("3.  Draft License Agreement — salient terms incorporated")
p("In compliance with Resolution 4(b), the draft at **Part-B** has been prepared as a **licence** (and not a lease/tenancy) so that no interest in the Institute’s property is created and the permission remains revocable in accordance with the Agreement. The monthly amount resolved as “rent” by the CEC is described in the draft as **“License Fee”**, for the reasons recorded above. The CEC-approved financial terms have been incorporated **verbatim, without any change**:")
table([
    ["Term", "CEC Resolution (Item 4)", "Clause in Draft (Part-B)"],
    ["License Fee / Rent", "Rs. 3,000/- per month", "Clause 4.1"],
    ["Security Deposit", "Rs. 30,000/-", "Clause 5.1"],
    ["Term", "One (1) year", "Clause 3.1"],
    ["Water Charges", "Standard rates until water meter is installed", "Clause 6.2 — thereafter metered consumption"],
    ["Executing an agreement & forwarding draft to Society", "Resolution 4(b)", "Part-B (execution) & Part-C (forwarding letter)"],
], widths=[2.0, 2.6, 1.97], size=10)

h2("4.  Financial implication (at approved rates)")
table([
    ["Head", "Amount / Basis"],
    ["License Fee", "Rs. 3,000/- × 12 months = **Rs. 36,000/- per annum** [[+ GST / other taxes, if applicable]]"],
    ["Security Deposit (refundable)", "Rs. 30,000/- (one-time, interest-free)"],
    ["Water Charges", "At standard Institute rates until meter installation; actuals thereafter"],
    ["Electricity Charges", "[[As per Clause 6.1 — sub-meter actuals / standard rates: for specific approval]]"],
    ["Other", "Taxes, statutory levies and stamp duty as applicable — Clauses 4, 7 & 23.8"],
], widths=[2.2, 4.37], size=10)

h2("5.  Salient safeguards built into the draft")
bullet("Premises to be used **only** for the coffee machine / approved beverages & items at CEC-approved prices displayed at the outlet (Clauses 8–9; Schedule-II). No cooking with fire; only the permitted machine/heating as approved.")
bullet("**FSSAI** licence/registration, sourcing only from FSSAI-registered suppliers, quality and hygiene responsibility on the Licensee — consistent with the spirit of CEC Resolutions under Item Nos. 1–2 (Clause 10).")
bullet("**Library discipline** — operation restricted to the earmarked café area and Library hours; no disturbance to readers; compliance with directions of the Librarian and the Professor-in-Charge (Commercial Establishment) (Clause 11).")
bullet("No assignment, sub-letting or parting with possession; no structural alteration; cleanliness, waste disposal, safety and personnel-verification obligations (Clauses 12–16).")
bullet("Termination for default / breach and for the Institute’s requirement; vacation and restoration on expiry/termination; penal charge for unauthorised overstay (Clauses 18–19).")

h2("6.  Points specifically submitted for approval / confirmation")
p("The following matters are **not expressly covered in the CEC Minutes** and have been proposed on the basis of standard practice. The Competent Authority is requested to approve the same or indicate modifications (the corresponding clauses will be conformed before execution):")
table([
    ["Sl.", "Matter (proposed in draft)", "Clause"],
    ["(a)", "Characterisation as **licence fee** (CEC term: “rent”); payable **in advance by the 7th day** of each month; interest on delayed payment [[@ 1% per month — or as approved]]", "4"],
    ["(b)", "License Fee [[exclusive of GST and other taxes, if/applicable at prevailing rates]]; TDS / tax compliance by Licensee", "4 & 7"],
    ["(c)", "**Electricity charges** — on sub-meter actuals where provided, otherwise at standard Institute rates (Minutes are silent on electricity)", "6.1"],
    ["(d)", "Operating hours to align with **Library hours** [[presently ______ hrs — to be confirmed by Librarian]] and list of approved items & prices in Schedule-II to be approved before commencement", "9 & 11"],
    ["(e)", "Termination for Institute’s requirement / convenience on **30 days’ notice**; immediate termination for grave / food-safety defaults", "18"],
    ["(f)", "Penal charges for overstay [[at three times the monthly License Fee, pro-rata — or as approved]]", "19.4"],
    ["(g)", "No automatic renewal; renewal only with **fresh approval of CEC / Competent Authority** and on revised terms, if any", "3"],
    ["(h)", "Value of non-judicial stamp paper (Karnataka Stamp Act) and legal vetting, if required, before execution — stamp duty to be borne by the Licensee", "23.8"],
    ["(i)", "Authorisation to the **Registrar** to execute the Agreement on behalf of the Institute, and delegation for day-to-day directions to PIC (Commercial Establishment) / Librarian", "—"],
], widths=[0.45, 4.82, 1.3], size=9.5)
p("Blanks highlighted in **yellow** throughout Part-B (area & plan of the premises, Society’s registration & authorised signatory, commencement/expiry dates, payment account details, operating hours, approved price list, etc. — consolidated checklist at Para 8 below) will be filled in, after joint verification, **before** execution. No blank will be left in the executed Agreement.", base_size=10)

h2("7.  Decision sought")
numbered("(a)", "**Approval of the Draft License Agreement at Part-B**, including the points at Para 6 above (with modifications, if any, as may be directed);")
numbered("(b)", "Approval to **forward the approved draft to the NITK Consumer Co-operative Society Ltd.** under the draft letter at **Part-C**, in terms of CEC Resolution 4(b), calling upon the Society to convey acceptance, deposit Rs. 30,000/-, and execute the Agreement within the time stipulated therein; and")
numbered("(c)", "Authorisation to the **Registrar, NITK Surathkal**, to execute the Agreement on behalf of the Institute after completion of blanks, payment of security deposit and verification of the Society’s FSSAI licence and authorisation of its signatory.")

h2("8.  Checklist of blanks / documents before execution")
table([
    ["Sl.", "Item to be filled / obtained", "To be furnished / confirmed by", "Status"],
    ["1", "Earmarked area (sq. ft.), floor/room & sketch plan of Akshara Café space — Schedule-I", "Librarian / PIC (Civil Infra.), jointly with Society", "[[ ]]"],
    ["2", "Commencement date (date of handing over) & expiry date (1 year)", "PIC (Commercial Estb.)", "[[ ]]"],
    ["3", "Society’s Registration No., registered address, PAN & GST (if any)", "Society", "[[ ]]"],
    ["4", "Certified copy of Society’s resolution authorising the signatory (name & designation)", "Society", "[[ ]]"],
    ["5", "Security Deposit of Rs. 30,000/- — DD / NEFT reference & date (Clause 5)", "Society / Accounts", "[[ ]]"],
    ["6", "FSSAI licence / registration copy; supplier details (Clause 10)", "Society", "[[ ]]"],
    ["7", "Approved items & price list — Schedule-II; operating hours (Clauses 9 & 11)", "CEC / Competent Authority; Librarian", "[[ ]]"],
    ["8", "Water (standard rate) & electricity charging basis till meters are provided — Clause 6", "PIC (Civil / Electrical Infra.)", "[[ ]]"],
    ["9", "Bank account / payment details of the Institute for monthly License Fee — Clause 4.3", "Accounts Section-I", "[[ ]]"],
    ["10", "Stamp paper value & purchase in the name of [[Licensee]]; legal vetting, if directed", "Accounts / Legal", "[[ ]]"],
    ["11", "Joint handing-over inventory — Schedule-III", "Librarian / PIC & Society", "[[ ]]"],
], widths=[0.4, 3.1, 2.05, 1.02], size=9.5)

h2("9.  Recommendation")
p("Submitted for kind perusal and approval of the Competent Authority. On approval, Accounts Section-I will (i) fair the Agreement incorporating the modifications, if any, as approved; (ii) forward the draft to the Society under Part-C; and (iii) place the executed Agreement on the respective file, with copies to the Librarian, PIC (Commercial Establishment) and PICs (Civil & Electrical Infrastructure) for metering, billing and monitoring.")

p("", space_after=6)
sig_table([
    (["**Prepared / Submitted by:**", "", "Signature: ___________________________",
      "Name: [[______________________________]]", "Assistant Registrar (Accounts)", "Date: [[______________]]"],
     ["**Recommended by:**", "", "Signature: ___________________________",
      "Name: [[______________________________]]", "Professor-in-Charge (Commercial Establishment)", "Date: [[______________]]"]),
    (["", "Signature: ___________________________",
      "Name: [[______________________________]]", "Deputy Registrar (Finance)", "Date: [[______________]]"],
     ["", "Signature: ___________________________",
      "Name: [[______________________________]]", "Registrar, NITK Surathkal", "Date: [[______________]]"]),
    (["**Recommended / Forwarded:**", "", "Signature: ___________________________",
      "Name: [[______________________________]]", "Dean (FSW) — Chairman, CEC", "Date: [[______________]]"],
     ["**Decision of the Competent Authority:**", "[[ ☐ Approved    ☐ Approved with modifications (noted)    ☐ Not approved ]]",
      "", "Signature: ___________________________", "**Director, NITK Surathkal**", "Date: [[______________]]"]),
])
p("*Modification(s), if any, directed at the time of approval may be recorded here / on the draft, and the fair draft will be conformed accordingly:*  [[______________________________________________________________________________]]", base_size=9.5)

# ============================================================
# PART B — DRAFT LICENSE AGREEMENT
# ============================================================
pagebreak()
p_center("PART – B", size=12, bold=True, color=MAROON, space_after=1)
p_center("Draft License Agreement (Execution Text)", size=13, bold=True, color=NAVY, space_after=2)
# stamp paper note box
t3 = doc.add_table(rows=1, cols=1)
t3.style = 'Table Grid'; t3.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = t3.cell(0, 0); shade(cell, LIGHT_GREY_BG); cell.paragraphs[0].clear()
para = cell.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = para.add_run("[ To be engrossed on Non-Judicial Stamp Paper of Rs. __________ purchased in the State of Karnataka, as applicable under the Karnataka Stamp Act, 1957 ]")
r.font.name = 'Times New Roman'; r.font.size = Pt(9); r.italic = True
doc.add_paragraph().paragraph_format.space_after = Pt(6)

h1("LICENSE AGREEMENT", center=True)
p_center("इज़ाज़त (लाइसेंस) अनुबंध", size=11, color=GREY, space_after=6)

p("THIS **LICENSE AGREEMENT** (hereinafter referred to as the **“Agreement”**) is made and executed on this [[____]] day of [[____________]], "
  "[[2026]] at Surathkal, Mangaluru,")
p_center("BY AND BETWEEN", size=11, bold=True, space_before=6, space_after=4)
p("**NATIONAL INSTITUTE OF TECHNOLOGY KARNATAKA, SURATHKAL**, an Institute of National Importance established under the National Institutes of Technology, "
  "Science Education and Research Act, 2007 (NITSER Act), Government of India, having its campus and office at **NITK Surathkal, Srinivasnagar Post, Mangaluru – 575 025, "
  "Dakshina Kannada District, Karnataka**, represented herein by its **Registrar** [[, (Name: ____________________________)]] , duly authorised in this behalf "
  "(hereinafter referred to as the **“Institute”** or the **“Licensor”**, which expression shall, unless repugnant to the context, include its successors and permitted assigns) "
  "of the **FIRST PART**;")
p_center("AND", size=11, bold=True, space_before=6, space_after=4)
p("**NITK CONSUMER CO-OPERATIVE SOCIETY LTD.**, a co-operative society registered under [[the Karnataka Co-operative Societies Act, 1959 / "
  "______________________________]] , bearing Registration No. [[____________________]] dated [[______________]], having its registered office at "
  "[[____________________________________________, NITK Campus, Surathkal, Mangaluru – 575 025]] , represented herein by its "
  "[[President / Secretary]] , [[(Name: ____________________________, aged ____, S/o/D/o ____________________________)]] , duly authorised vide resolution of its "
  "Managing Committee / General Body dated [[______________]] , a certified copy whereof is enclosed with this Agreement (hereinafter referred to as the "
  "**“Society”** or the **“Licensee”**, which expression shall, unless repugnant to the context, include its successors and permitted assigns) of the **SECOND PART**.")
p("The Licensor and the Licensee are individually referred to as a **“Party”** and collectively as the **“Parties”**.")

h3("WHEREAS:")
numbered("A.", "The Institute is the owner and is in possession of its campus at Surathkal, Mangaluru, including the Central Library / E-Library building and the earmarked space therein for a café facility for the students, faculty, staff and bonafide visitors of the Institute, known as **“Akshara Café”** (more particularly described in **Schedule-I** hereto and hereinafter referred to as the **“Licensed Premises”**).")
numbered("B.", "The Licensee requested the Institute for permission to install and operate a **coffee machine** at Akshara Café in the E-Library premises for serving coffee and other approved beverages/items to the users of the Library and the Institute community.")
numbered("C.", "The **Commercial Establishment Committee (CEC)** of the Institute, in its meeting held on **23-09-2026**, vide **Item No. 4** of the Minutes issued under Ref. No. **07/NITK/AC-1/2026-27/COM-ESTB/A6 dated 24-09-2026**, approved the said request on the terms, inter alia, of monthly rent of Rs. 3,000/-, security deposit of Rs. 30,000/-, contract period of one (1) year and water charges at standard rates until a water meter is installed, and directed that the Licensee shall sign an agreement with the Institute.")
numbered("D.", "In pursuance of the aforesaid Resolution, the Institute has agreed to **grant, and the Licensee has agreed to accept, a bare, non-exclusive and revocable licence** to use the Licensed Premises, upon and subject to the terms and conditions herein set forth, and **not** by way of lease, tenancy or any interest in the property.")
numbered("E.", "The Licensee has inspected the Licensed Premises, is satisfied as to its suitability for the permitted use, and agrees to accept it on an **“as-is-where-is”** basis.")

p("**NOW THEREFORE**, in consideration of the License Fee and other charges payable hereunder and of the covenants herein, the Parties agree as follows:", space_before=6)

# -- Clause 1
h2("1.  Definitions and Interpretation")
subclause("1.1", "“**Licensed Premises**” means the earmarked space for Akshara Café in the E-Library premises described in Schedule-I, together with the fixtures, if any, listed in Schedule-III, and no other area.")
subclause("1.2", "“**License Fee**” means the monthly fee payable under Clause 4 (referred to as “rent” in the CEC Minutes dated 23-09-2026).")
subclause("1.3", "“**Commencement Date**” means the date on which possession of the Licensed Premises is handed over to the Licensee for the permitted use, as recorded in Schedule-I/III, or [[the seventh day from execution of this Agreement]], whichever is earlier.")
subclause("1.4", "“**Competent Authority**” means the Director, NITK Surathkal, or the authority to whom powers are delegated in this behalf; “**PIC**” means the Professor-in-Charge (Commercial Establishment) or any other officer nominated by the Institute.")
subclause("1.5", "Headings are for convenience only. The Schedules and Annexure-A form part of this Agreement. References to statutes include amendments and re-enactments in force.")

h2("2.  Grant of Licence and Nature of Rights")
subclause("2.1", "The Licensor hereby grants to the Licensee a **bare, non-exclusive, non-transferable and revocable licence** to enter upon and use the Licensed Premises solely for the permitted use stated in Clause 8, during the Term (Clause 3), subject to compliance with this Agreement.")
subclause("2.2", "This Agreement is a **licence within the meaning of Section 52 of the Indian Easements Act, 1882**. It does **not** create, and shall not be deemed to create, any tenancy, lease, easement, possession in law, or any right, title or interest whatsoever in the Licensed Premises or any part of the Institute’s property. Legal possession and control of the Licensed Premises shall at all times remain with the Institute.")
subclause("2.3", "The Licensee shall have no right to exclude the Institute, its officers, or persons authorised by it, from entering the Licensed Premises at any reasonable time (and at any time in an emergency) for inspection, maintenance, safety or any lawful purpose.")
subclause("2.4", "Nothing herein shall oblige the Institute to provide any furniture, equipment, manpower or service beyond handing over the earmarked space and the existing utility points, if any, recorded in Schedule-III. The coffee machine, accessories, consumables and personnel shall be arranged by the Licensee entirely at its cost and risk.")

h2("3.  Term and Renewal")
subclause("3.1", "The licence shall commence on the **Commencement Date**, i.e., [[____-____-______]] , and shall remain in force for a period of **one (1) year**, expiring on [[____-____-______]] (both days inclusive), unless terminated earlier in accordance with Clause 18 (the “**Term**”). Time shall be of the essence.")
subclause("3.2", "There shall be **no automatic renewal or extension**. Any renewal shall require the **prior written approval of the CEC / Competent Authority**, a fresh or renewed agreement, and such revised license fee, security deposit and other terms as may then be resolved. The Licensee shall apply for renewal, if desired, at least [[sixty (60) days]] before expiry; continued occupation after expiry without such written approval shall be treated as unauthorised (Clause 19).")
subclause("3.3", "During vacation, closure of the Library/Institute, force majeure or for any other reason, no proportionate reduction, suspension or waiver of the License Fee or other charges shall be claimed, except as may be expressly approved in writing by the Competent Authority.")

h2("4.  License Fee, Payment and Delay")
subclause("4.1", "The Licensee shall pay to the Institute a monthly **License Fee of Rs. 3,000/- (Rupees Three Thousand only)** for the use of the Licensed Premises during the Term.")
subclause("4.2", "The License Fee for every calendar month shall be paid **in advance, on or before the seventh (7th) day** of that month; for a broken period it shall be paid pro-rata before commencement. [[The first month’s License Fee shall be paid before / at the time of handing over.]]")
subclause("4.3", "All payments shall be made by [[account-payee demand draft drawn in favour of “The Director, National Institute of Technology Karnataka, Surathkal”, payable at Mangaluru / NEFT–RTGS to Account No. __________, IFSC __________, Bank & Branch ____________________, as intimated by Accounts Section-I in writing]] , and the Licensee shall share the payment reference promptly to obtain the Institute’s receipt. Cash payment shall not be made unless expressly permitted.")
subclause("4.4", "The License Fee is [[exclusive of Goods and Services Tax (GST) and all other taxes, duties and levies, which, if and as applicable, shall be charged and paid additionally at the prevailing rates]] . Tax deducted at source, if applicable, shall be deposited by the Licensee and TDS certificate issued to the Institute in time.")
subclause("4.5", "Delay in payment of the License Fee or any other sum due beyond the due date shall attract [[interest / late fee @ 1% (one per cent) per month (or such rate as may be approved / notified by the Institute from time to time)]] on the amount outstanding for the period of delay, without prejudice to the Institute’s right to terminate under Clause 18. Default in payment for [[two consecutive months]] shall, by itself, be a ground for termination.")
subclause("4.6", "No dispute, claim or set-off of any nature shall entitle the Licensee to withhold or defer payment of the License Fee, utility charges or any other amount due under this Agreement.")
subclause("4.7", "Upon renewal, if approved, the License Fee and other charges shall be such as are resolved by the CEC / Competent Authority for the renewed term.")

h2("5.  Security Deposit")
subclause("5.1", "Before or at the time of execution/handling over, the Licensee shall deposit with the Institute an **interest-free refundable Security Deposit of Rs. 30,000/- (Rupees Thirty Thousand only)** by [[demand draft / NEFT, as per Clause 4.3 — DD No. / UTR No. ____________ dated ____________ for Rs. 30,000/- is acknowledged]] , as security for due performance of this Agreement.")
subclause("5.2", "The Security Deposit shall not be treated as advance License Fee and shall not carry any interest. The Institute may, without prejudice to its other rights, adjust / appropriate from it any License Fee, utility charges, taxes, damages, costs of repairs or restoration, penal charges or other sums due from the Licensee, after giving intimation.")
subclause("5.3", "If any part of the Security Deposit is so adjusted during the Term, the Licensee shall replenish it to the full amount within [[seven (7) days]] of demand.")
subclause("5.4", "Subject to adjustments under Clause 5.2 and to the Licensee handing over vacant possession and restoring the Licensed Premises as required by Clause 19, the Security Deposit (or balance thereof) shall be refunded within [[thirty (30) days]] of expiry or earlier termination and clearance of all dues, without interest.")

h2("6.  Water, Electricity and Other Utility Charges")
subclause("6.1", "**Electricity:** The Licensee shall bear the cost of electricity consumed for the coffee machine and other permitted equipment. [[Where a sub-meter is provided, charges shall be on actual consumption at the Institute’s applicable tariff; till such time / where no sub-meter is provided, charges shall be levied at the standard rates / on the basis of connected load as assessed by the PIC (Electrical Infrastructure), which assessment shall be final.]]")
subclause("6.2", "**Water:** As resolved by the CEC, **water charges shall be levied at standard rates until a water meter is installed** for the Licensed Premises. Upon installation of a water meter, charges shall be recovered on the basis of **actual metered consumption** at the rates applicable to commercial establishments of the Institute from time to time. The Licensee shall afford full access and cooperation for installation, reading, sealing, maintenance and replacement of meters and shall not tamper with them in any manner.")
subclause("6.3", "Meter readings recorded jointly or by the Institute’s authorised representative, and standard-rate assessments made during the pre-meter period, shall be final and binding. The Licensee shall pay water, electricity and any other utility/service charges (including drainage, waste disposal or common services, if levied) in addition to the License Fee, **within** [[seven (7) days]] **of the bill / demand**; Clause 4.5 shall apply to delayed payment.")
subclause("6.4", "The Institute shall not be liable for any interruption, fluctuation or failure of water, power or other services, nor for any loss or damage arising therefrom. The Licensee shall make its own arrangements, if permitted in writing, at its cost, without causing danger or nuisance.")
subclause("6.5", "The Licensee shall use water and electricity economically, prevent wastage and leakage, and promptly report any defect to the PIC / Librarian and the concerned maintenance section.")

h2("7.  Taxes and Statutory Levies")
subclause("7.1", "All taxes, duties, cesses, licence fees and other statutory levies in respect of the Licensee’s business, equipment, sales and personnel shall be borne and paid by the Licensee alone. The Institute’s property tax, if any, in respect of the building shall remain the Institute’s responsibility, unless a levy is specifically attributable to the Licensee’s commercial use, in which case it shall be reimbursed by the Licensee on demand.")
subclause("7.2", "The Licensee shall obtain and keep valid, at its cost, all permissions, registrations and licences required under applicable law, including under the Food Safety and Standards Act, 2006, the Karnataka Shops and Commercial Establishments law (if applicable), labour, and municipal / local authority requirements, and produce copies to the Institute before commencement and upon every renewal.")

h2("8.  Permitted Use and Prohibitions")
subclause("8.1", "The Licensed Premises shall be used **solely for installing and operating a coffee machine and for the sale/service of coffee, tea and such other beverages and items as are expressly approved in writing** by the CEC / Competent Authority and listed in **Schedule-II**, at “Akshara Café”, to students, faculty, staff and bonafide visitors of the Institute. **No other trade, business, storage or activity** shall be carried on without prior written approval.")
subclause("8.2", "**No cooking of any kind** involving open fire, gas, kerosene or other hazardous / inflammable material is permitted in or around the Licensed Premises. Only the permitted coffee machine and such heating of ready-made approved items [[by microwave oven, if expressly approved in Schedule-II]] shall be allowed, in conformity with fire-safety and Library requirements.")
subclause("8.3", "The sale, storage or consumption of **alcohol, tobacco and tobacco products, gutkha, narcotic / psychotropic substances, and any other item prohibited by law or by the Institute** is strictly forbidden. [[Non-vegetarian items shall be sold only if expressly approved in Schedule-II.]] Single-use plastic shall not be used contrary to applicable rules / Institute directions.")
subclause("8.4", "The Licensee shall not use any space outside the Licensed Premises (including corridors, reading halls, staircases or open areas) for service, seating, display, storage or advertisement, except seating, if any, expressly earmarked in Schedule-I. No encroachment or obstruction shall be caused.")
subclause("8.5", "No name board, banner, advertisement or branding shall be displayed except a board of the size, design and location approved in writing by the Institute / Librarian. The name “**Akshara Café**” and the Institute’s name / emblem shall not be used for any outside business or publicity without prior written approval.")

h2("9.  Items, Prices and Display")
subclause("9.1", "Only items listed in **Schedule-II** shall be sold, and **strictly at the prices and in the quantities / measures approved** therein by the CEC / Competent Authority. The Licensee shall not charge any amount over and above the approved price / MRP of packaged items, as applicable.")
subclause("9.2", "The approved price list, FSSAI licence/registration, and the name & contact of the person-in-charge (with a notice inviting complaints to the PIC (Commercial Establishment) [[/ Librarian]]) shall be conspicuously displayed at the outlet at all times, and bills/receipts shall be issued on demand; [[digital payment facility shall be provided]].")
subclause("9.3", "Any addition of items or revision of prices / quantities shall require **prior written approval** of the CEC / Competent Authority. [[Till Schedule-II is approved and signed, the Licensee shall not commence sale of any item other than by the coffee machine at the prices communicated in writing.]]")

h2("10.  Food Safety, Quality, Hygiene and Sourcing")
subclause("10.1", "The Licensee shall, **before commencement**, obtain a valid **FSSAI licence / registration** for the outlet / activity, bearing No. [[____________________ valid up to ____________]] , keep it subsisting during the entire Term, display it as required, and comply with the Food Safety and Standards Act, 2006 and the Rules/Regulations thereunder, and all guidelines of FSSAI and the Institute.")
subclause("10.2", "All beverages, ingredients and items shall be fresh, wholesome and of good quality, procured **only from FSSAI-registered / licensed suppliers**. The Licensee shall furnish, before commencement and whenever required or changed, the details of its sources/suppliers, and shall be **solely responsible for the quality, safety and fitness for consumption of everything sold/served**. Expired, stale, adulterated or unhygienically stored items shall not be stocked or sold.")
subclause("10.3", "Only potable water shall be used for preparation of beverages. The coffee machine, utensils, dispensers and storage shall be cleaned and, where required, sterilised daily; food-contact surfaces shall be kept hygienic; and pest control shall be carried out regularly at the Licensee’s cost without endangering Library materials or users.")
subclause("10.4", "Personnel handling beverages shall be medically fit, free from communicable diseases, maintain personal hygiene, and wear clean [[uniform / apron, cap and gloves]] while on duty; medical fitness certificates shall be produced on demand.")
subclause("10.5", "The Institute, the CEC, the PIC, the Librarian, the Institute’s medical / food-safety personnel, or any statutory authority may inspect and take samples at any time. Any deficiency pointed out shall be rectified forthwith. Sale of unsafe or substandard items, or any incident of food poisoning attributable to the Licensee, shall entitle the Institute to immediately suspend operations and/or terminate this Agreement under Clause 18, apart from other legal remedies, and the Licensee shall bear all consequences thereof.")

h2("11.  Operation, Hours and Library Discipline")
subclause("11.1", "The café shall operate only during the hours approved by the Librarian / Institute, presently [[from ______ hrs to ______ hrs on ____________ (days), aligned to E-Library hours]] , which may be revised in writing. The outlet shall remain closed when the Library is closed, except with express written permission.")
subclause("11.2", "The Licensee understands that the Licensed Premises are within the Library premises. Silence, decorum and the academic atmosphere shall be strictly maintained; no noise, music, crowding, loitering or activity disturbing readers or Library functioning shall be permitted. The directions of the **Librarian** in all matters concerning Library discipline, entry, security and use of the premises shall be final and binding.")
subclause("11.3", "Beverages / items shall be consumed **only within the earmarked café area**. The Licensee shall not serve, and shall use reasonable efforts to discourage carrying of, beverages/food into reading halls, stack rooms or other prohibited areas of the Library.")
subclause("11.4", "Entry of the Licensee, its personnel, suppliers and customers shall be subject to the security and Library rules of the Institute, including identity checks and material gate-pass procedures. Deliveries shall be made at the times and through the routes approved by the Librarian / Security, without disturbance.")
subclause("11.5", "The service shall be run efficiently and continuously during approved hours on all approved working days. Unauthorised closure for more than [[three (3) consecutive days / seven (7) days in a month]] without prior written permission shall be treated as a breach.")
subclause("11.6", "Complaints, if any, from users shall be attended to promptly and courteously. A complaint / suggestion register shall be maintained at the outlet, produced to the PIC / Librarian on demand, and complaints shall be reported as directed.")

h2("12.  Personnel")
subclause("12.1", "All persons engaged by the Licensee shall be its own personnel. The list of personnel (with photographs, identity and address proof and police verification, where required by the Institute) shall be furnished before deployment and upon every change; they shall wear Institute-approved identity cards while on campus.")
subclause("12.2", "The Licensee shall ensure that its personnel are well-behaved, courteous and disciplined, do not indulge in smoking, consumption of alcohol or prohibited substances, or any misconduct, and comply with all campus and Library rules. Any person objected to by the Institute shall be replaced forthwith.")
subclause("12.3", "There shall be **no employer–employee, master–servant or principal–agent relationship** between the Institute and the Licensee or its personnel. The Licensee alone shall be responsible for salaries, wages, EPF, ESI, bonus, leave, compensation and all statutory obligations towards its personnel, and shall indemnify the Institute in that behalf. No personnel of the Licensee shall have or claim any right of employment, absorption or regularisation in the Institute.")
subclause("12.4", "No personnel (or family member) of the Licensee shall reside or stay overnight in the Licensed Premises or anywhere on the campus under this Agreement.")

h2("13.  Maintenance, Cleanliness, Waste and Safety")
subclause("13.1", "The Licensee shall keep the Licensed Premises, the machine and the surroundings used by it **clean, neat, odour-free and in good order** at its own cost, to the satisfaction of the Institute / Librarian. Daily cleaning, including removal of spills that may damage flooring, furniture or Library property, shall be the Licensee’s responsibility.")
subclause("13.2", "All waste (including used cups, grounds and liquid waste) shall be segregated and disposed of daily by the Licensee at its cost through the Institute-approved system / location only, in compliance with municipal and environmental norms. Liquid waste shall not be discharged so as to block drains or create nuisance. No littering shall be permitted.")
subclause("13.3", "The Licensee shall comply with all fire-safety, electrical-safety and general safety norms. The coffee machine and electrical connections shall be of approved standard, properly earthed and maintained by competent persons; no overloading, unauthorised tapping or unsafe wiring shall be done. [[A suitable fire extinguisher of approved type shall be kept at the outlet, if directed.]]")
subclause("13.4", "The Licensee shall be responsible for the safety and security of its machine, equipment, stock and cash. The Institute shall not be liable for any theft, loss or damage thereto from any cause.")
subclause("13.5", "Any damage caused by the Licensee, its personnel, suppliers or customers to the building, fittings, furniture, books/materials or any Institute property shall be made good / compensated by the Licensee at the assessment of the Institute, recoverable also from the Security Deposit.")

h2("14.  Alterations, Fixtures and Equipment")
subclause("14.1", "No structural addition, alteration, demolition, construction, partitioning, drilling into or painting of walls/floors, or change in electrical/plumbing lines shall be made without the **prior written approval** of the Institute (through the PICs — Civil and Electrical Infrastructure). All approved works shall be executed at the Licensee’s cost as per the Institute’s specifications and supervision.")
subclause("14.2", "The coffee machine and all other equipment, counters and movable articles brought in by the Licensee shall be recorded in **Schedule-III**, shall remain its property, and shall be removed by it on expiry/termination, making good all damage caused by installation or removal. Fixtures permanently annexed with approval shall, unless otherwise directed in writing, be removed and the premises restored at the Licensee’s cost.")
subclause("14.3", "The Institute’s fixtures and articles, if any, handed over with the premises (Schedule-III) shall be maintained in good condition and returned in the same condition, fair wear and tear excepted; loss or damage shall be compensated.")

h2("15.  Inspection and Compliance with Directions")
subclause("15.1", "The Institute, through the PIC (Commercial Establishment), the Librarian, the Registrar or any authorised officer/committee (including the CEC), may inspect the Licensed Premises, machine, stock, records, bills and price display at any time without notice. The Licensee shall give full access, produce records and afford all cooperation.")
subclause("15.2", "The Licensee shall strictly comply with the NITSER Act, the Institute’s Statutes, Rules and Regulations, the resolutions/directions of the CEC, and all written instructions of the Competent Authority, the PIC and the Librarian in relation to the running of the café, prices, items, hygiene, discipline and campus security, issued from time to time. In case of any inconsistency, the Institute’s direction shall prevail.")

h2("16.  No Assignment, Sub-letting or Parting with Possession")
subclause("16.1", "The licence is **personal to the NITK Consumer Co-operative Society Ltd.** The Licensee shall not assign, transfer, sub-let, sub-license, franchise, mortgage or in any manner part with possession or allow use of the Licensed Premises or any right under this Agreement, in whole or in part, to any other person, firm, society or agency, nor induct any partner/operator to run the outlet. Any such act shall be **null and void** and shall be a ground for immediate termination.")
subclause("16.2", "A change in the Managing Committee / office bearers of the Society in the normal course shall be intimated to the Institute within [[fifteen (15) days]] with certified copies, and shall not, by itself, be treated as assignment if the Society continues as the same legal entity and the Institute raises no objection.")

h2("17.  Insurance, Indemnity and Liability")
subclause("17.1", "The Licensee may, at its cost, insure its machine, equipment and stock and take such third-party / public liability insurance as it considers prudent [[/ as may be directed by the Institute]]; failure to insure shall not fasten any liability upon the Institute.")
subclause("17.2", "The Licensee shall **indemnify and keep indemnified** the Institute, its officers, faculty, staff and students against all claims, demands, actions, losses, damages, penalties and expenses arising out of or in connection with the Licensee’s operations, including those relating to food safety/quality, its personnel, suppliers, equipment, breach of law, or breach of this Agreement.")
subclause("17.3", "The Institute shall not be liable, directly or indirectly, for any loss, damage, injury or claim arising from the consumption of any item sold, the operation of the machine, interruption of services, or any act or omission of the Licensee, its personnel or suppliers.")

h2("18.  Default, Suspension and Termination")
subclause("18.1", "**Termination for breach / default.** Without prejudice to any other remedy, the Institute may terminate this Agreement forthwith (or by giving such shorter notice / opportunity to cure as it deems fit, having regard to the nature of the breach) if the Licensee:")
for lbl, txt in [
    ("(a)", "fails to pay the License Fee, utility charges or any other sum by the due date, or defaults as in Clause 4.5;"),
    ("(b)", "fails to deposit or replenish the Security Deposit;"),
    ("(c)", "sells unapproved items, charges more than the approved price, or indulges in unfair trade practice;"),
    ("(d)", "violates food-safety, quality or hygiene requirements, or its FSSAI licence is suspended/cancelled or lapses;"),
    ("(e)", "sublets/assigns or parts with possession, makes unauthorised alterations, or uses the premises for a non-permitted purpose;"),
    ("(f)", "violates Library/campus discipline, safety norms or any law, or fails to comply with the directions of the Institute / CEC / Librarian;"),
    ("(g)", "closes the outlet without permission beyond the period in Clause 11.5, or abandons the premises;"),
    ("(h)", "is dissolved, becomes insolvent, or ceases to function / is superseded in a manner affecting performance; or"),
    ("(i)", "commits any other breach of this Agreement."),
]:
    subclause("    " + lbl, txt)
subclause("18.2", "**Termination for Institute’s requirement / convenience.** The Institute may revoke the licence and terminate this Agreement at any time, without assigning any reason, by giving **thirty (30) days’ written notice**, including where the premises or the Library space is required by the Institute for its own purposes, reorganisation, repair or any policy decision of the CEC / Competent Authority. In such case, dues shall be settled pro-rata and the Security Deposit dealt with under Clause 5.4; no compensation for loss of business or goodwill shall be payable.")
subclause("18.3", "**Termination by the Licensee.** The Licensee may surrender the licence by giving **thirty (30) days’ written notice** [[/ payment of License Fee for the notice period in lieu of notice]], subject to clearing all dues and fulfilling Clause 19. [[Surrender within the first ______ months, if so approved, may entail forfeiture as directed by the Competent Authority.]]")
subclause("18.4", "**Suspension.** Pending inquiry into a complaint, food-safety incident, safety hazard or grave breach, the Institute may direct immediate suspension of operations, and the Licensee shall comply forthwith, without any claim for waiver or compensation.")
subclause("18.5", "For any breach not resulting in termination, the Institute may, instead of or in addition to other action, impose a penalty [[up to Rs. ____________ per breach / per day of continuing breach, as assessed by the PIC and approved by the Competent Authority]], recoverable from amounts due / the Security Deposit, after giving the Licensee an opportunity to explain.")

h2("19.  Consequences of Expiry or Termination; Handing Over")
subclause("19.1", "On expiry of the Term or on termination, the licence shall **automatically stand revoked**. The Licensee shall, **on that very date**, peacefully quit and hand over vacant and clean possession of the Licensed Premises to the Institute (through the Librarian / PIC), remove its machine, equipment and stock, restore the premises to the condition in which they were handed over (fair wear and tear excepted), and settle all dues.")
subclause("19.2", "A joint handing-over / taking-over memo, with meter readings and inventory verification (Schedule-III), shall be signed. The Licensee shall bear the cost of repairs/restoration for damage beyond fair wear and tear, as assessed by the Institute.")
subclause("19.3", "If the Licensee fails to remove its articles within [[seven (7) days]] after expiry/termination, the Institute may remove and store them at the Licensee’s risk and cost, and, if not claimed within [[thirty (30) days]], dispose of them in the manner it deems fit and adjust the proceeds/costs against dues, without liability.")
subclause("19.4", "**Unauthorised occupation.** For any period of overstay after expiry, termination or revocation, the Licensee shall be liable to pay **penal charges** [[at three times the monthly License Fee, on pro-rata basis]] for the period of unauthorised occupation, in addition to utility charges, taxes and damages, without prejudice to the Institute’s right to take legal, eviction and other proceedings at the Licensee’s cost. Acceptance of any amount shall not amount to renewal or waiver.")
subclause("19.5", "Termination shall not affect accrued rights and liabilities, including Clauses 5, 13.5, 17, 19, 21 and this Clause, which shall survive expiry or termination.")

h2("20.  Force Majeure")
subclause("20.1", "Neither Party shall be liable for failure to perform (other than payment obligations already accrued) if prevented by events beyond its reasonable control, including natural calamity, epidemic, Government order, or closure of the Institute by competent authority. The affected Party shall promptly intimate the other. If such event continues beyond [[thirty (30) days]] / substantially frustrates the Agreement, the Institute may terminate under Clause 18.2 mutatis mutandis.")
subclause("20.2", "Closure of the Library / Institute for vacations, examinations, maintenance or administrative reasons in the normal course shall not constitute force majeure for waiver of the License Fee (Clause 3.3), unless expressly approved in writing.")

h2("21.  Dispute Resolution and Jurisdiction")
subclause("21.1", "Any dispute, difference or question arising out of or in connection with this Agreement (including regarding quality, prices, meter readings, assessments, charges and penalties) shall first be referred to the **PIC (Commercial Establishment)** and, if not resolved, to the **Registrar / Chairman, CEC**. In matters of campus administration, Library discipline, approved items/prices and levy of charges in accordance with CEC resolutions, the decision of the **Director, NITK Surathkal** (or the CEC, as applicable) shall be final and binding on the Licensee.")
subclause("21.2", "Disputes not settled under Clause 21.1 may be resolved [[by arbitration by a sole arbitrator appointed by the Director, NITK Surathkal, in accordance with the Arbitration and Conciliation Act, 1996, with seat and venue at Mangaluru and language English]] — *and/or* [[to be conformed as approved / legally vetted]].")
subclause("21.3", "Subject to Clause 21.2, the Courts at **Mangaluru, Dakshina Kannada, Karnataka** shall have exclusive jurisdiction.")

h2("22.  Notices")
subclause("22.1", "All notices, bills and communications under this Agreement shall be in writing and shall be deemed duly served if delivered personally against acknowledgment, or sent by registered post / speed post / courier / e-mail (with confirmation) addressed as below, or to such other address as may be notified in writing:")
p("**For the Institute:** The Registrar, National Institute of Technology Karnataka, Surathkal, Mangaluru – 575 025; E-mail: [[____________________]] , with a copy to the Professor-in-Charge (Commercial Establishment) and the Librarian, NITK Surathkal.  **For the Licensee:** [[President / Secretary, NITK Consumer Co-operative Society Ltd., ____________________, Mangaluru – 575 025; Mobile: ____________ ; E-mail: ____________________]] .")
subclause("22.2", "A notice displayed at / pasted on the Licensed Premises in the event of refusal, evasion or closure, in the presence of two witnesses, shall also be deemed good service.")

h2("23.  General")
subclause("23.1", "**Entire Agreement.** This Agreement (with Schedules and Annexure-A) constitutes the entire understanding between the Parties in respect of the Licensed Premises and supersedes all prior oral or written understandings on the subject. The CEC Minutes dated 23-09-2026 form the basis of this Agreement; in case of any conflict on the financial terms resolved in Item No. 4 thereof, the Minutes shall prevail to the extent of such conflict, unless amended by the CEC / Competent Authority in writing.")
subclause("23.2", "**Amendment.** No amendment, waiver or relaxation of this Agreement shall be valid unless made in writing, approved by the CEC / Competent Authority and signed by both Parties.")
subclause("23.3", "**No waiver** by the Institute of any breach or of any provision shall operate as a waiver of any subsequent breach or of any other provision.")
subclause("23.4", "**Severability.** If any provision is held invalid or unenforceable, the remaining provisions shall continue in full force and effect.")
subclause("23.5", "**No partnership / agency.** Nothing herein creates any partnership, joint venture, agency or employment between the Parties. The Licensee shall not represent itself as an agent of the Institute, nor pledge the Institute’s credit.")
subclause("23.6", "The Licensee confirms that it has read and understood this Agreement, that its signatory is duly authorised, and that true copies of its registration, bye-laws (relevant extract), authorising resolution and FSSAI licence are enclosed / shall be furnished before commencement.")
subclause("23.7", "**Counterparts.** This Agreement may be executed in [[two]] counterparts, each of which shall be an original. The Parties shall initial every page and sign the Schedules.")
subclause("23.8", "**Stamp duty and costs.** Stamp duty and incidental charges for execution of this Agreement shall be borne by the **Licensee**. [[Each Party shall bear its own legal costs.]]")

p("**IN WITNESS WHEREOF** the Parties hereto have signed this Agreement on the day, month and year first hereinabove written.", space_before=8, space_after=8)

sig_table([
    (["**For and on behalf of the Licensor:**", "**NATIONAL INSTITUTE OF TECHNOLOGY KARNATAKA, SURATHKAL**",
      "", "Signature: ___________________________", "Name: [[______________________________]]",
      "Designation: **Registrar, NITK Surathkal**", "Date: [[______________]]   Place: Surathkal", "", "(Office Seal)"],
     ["**For and on behalf of the Licensee:**", "**NITK CONSUMER CO-OPERATIVE SOCIETY LTD.**",
      "", "Signature: ___________________________", "Name: [[______________________________]]",
      "Designation: [[President / Secretary]] — Authorised Signatory", "Date: [[______________]]   Place: [[______________]]", "", "(Seal of the Society)"]),
])

h3("Witnesses:")
sig_table([
    (["**1.**  Signature: ___________________________", "Name: [[______________________________]]",
      "Designation / Address: [[____________________________________________]]", "ID / Mobile: [[____________________]]"],
     ["**2.**  Signature: ___________________________", "Name: [[______________________________]]",
      "Designation / Address: [[____________________________________________]]", "ID / Mobile: [[____________________]]"]),
    (["**3.**  Signature: ___________________________  *(if any)*", "Name: [[______________________________]]",
      "Designation / Address: [[____________________________________________]]"],
     ["**4.**  Signature: ___________________________  *(if any)*", "Name: [[______________________________]]",
      "Designation / Address: [[____________________________________________]]"]),
])
p("*Enclosures to the executed Agreement: (i) CEC Minutes dated 23-09-2026 (Item No. 4 extract — Annexure-A); (ii) Society’s registration certificate & authorising resolution; "
  "(iii) FSSAI licence / registration; (iv) Society’s PAN (and GST, if any); (v) Security Deposit instrument/receipt; (vi) Plan of the Licensed Premises (with Schedule-I).*", base_size=9.5)

# ---- Schedule I
pagebreak()
h1("SCHEDULE – I", center=True)
p_center("Description of the Licensed Premises — “Akshara Café”, E-Library", size=11.5, bold=True, space_after=2)
p_center("(To be jointly measured, verified and signed before handing over; sketch plan to be attached and initialled by both Parties)", size=9.5, italic=True, space_after=8)
table([
    ["Sl.", "Particulars", "Details"],
    ["1", "Building / Location", "E-Library, [[Central Library Building, ______ Floor, Room / Bay No. ______]], NITK Campus, Surathkal, Mangaluru – 575 025"],
    ["2", "Earmarked area", "[[________ (length) × ________ (breadth) = ________ sq. ft. (approx.)]] — as shown in the attached sketch plan, marked in [[red]]"],
    ["3", "Boundaries / Identification", "North: [[__________]]   South: [[__________]]\nEast: [[__________]]   West: [[__________]]"],
    ["4", "Permitted installation", "One (1) coffee machine [[Make / Model: ____________________; Capacity: __________]] and its accessories, at the location marked in the plan"],
    ["5", "Seating (if earmarked)", "[[____ Nos. of chairs / tables within ________ area only / No seating earmarked]]"],
    ["6", "Utility points handed over", "Electrical point(s): [[__________]]   Sub-meter No.: [[__________]]  Initial reading: [[__________]]\nWater point(s): [[__________]]    Water meter No.: [[“Not installed — standard rates apply (Clause 6.2)” / __________]]  Initial reading: [[__________]]"],
    ["7", "Access", "Through [[____________________ entrance / route]], subject to Library and security rules (Clause 11)"],
    ["8", "Commencement / Handing-over Date", "[[____-____-______]] (Commencement Date, Clause 3.1)"],
    ["9", "Expiry Date (1 year)", "[[____-____-______]]"],
], widths=[0.45, 2.0, 4.12], size=10)
p("Attached: Sketch plan / layout of the Licensed Premises (signed): [[Plan No./Drawing reference ____________________ / Sheet attached at page _____]].")
p("The Licensee confirms that it has inspected the above premises, measured the area jointly with the representatives of the Librarian / Institute, and has taken over the same on an “as-is-where-is” basis for the permitted use.", space_after=8)
sig_table([
    (["**Handed over by (Institute):**", "", "Signature: ___________________________", "Name & Designation: [[____________________________]]",
      "([[Librarian / PIC (Commercial Estb.)]])", "Date & Time: [[________________]]"],
     ["**Taken over by (Licensee):**", "", "Signature: ___________________________", "Name & Designation: [[____________________________]]",
      "(For NITK Consumer Co-op. Society Ltd.)", "Date & Time: [[________________]]"]),
])

# ---- Schedule II
h1("SCHEDULE – II", center=True)
p_center("Approved Items and Prices at Akshara Café", size=11.5, bold=True, space_after=2)
p_center("(Only the items below, as approved in writing by the CEC / Competent Authority, shall be sold — Clauses 8 & 9. "
         "This Schedule shall be filled, approved and signed before commencement; any change requires prior written approval.)", size=9.5, italic=True, space_after=8)
rows = [["Sl. No.", "Item", "Quantity / Size (ml / g)", "Approved Price (Rs.)", "Remarks (source / brand, if specified)"]]
items = ["Coffee", "Tea", "Masala Tea", "[[Green Tea / Herbal Tea]]", "[[Hot Milk]]", "[[Badam / Flavoured Milk]]", "[[__________________]]", "[[__________________]]", "[[__________________]]", "[[__________________]]"]
for i, it in enumerate(items, start=1):
    rows.append([str(i), it, "[[______]]", "[[______]]", "[[____________________]]"])
table(rows, widths=[0.6, 2.1, 1.45, 1.25, 1.17], size=10,
      col_aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
p("**Conditions:**  (1) No item other than the above shall be sold without prior written approval.  (2) Packaged items, if any are approved, shall not be sold above MRP.  "
  "(3) All items shall be procured only from FSSAI-registered/licensed suppliers (Clause 10.2) — supplier details furnished: [[Name(s) & FSSAI No(s).: ________________________________________]].  "
  "(4) No cooking is permitted; only the permitted machine [[and microwave-heating of ready-made approved items, if expressly approved above]] is allowed (Clause 8.2).  "
  "(5) This price list shall be displayed at the outlet (Clause 9.2).", base_size=10)
p("Approved by (CEC / Competent Authority) vide [[______________ dated ______________]] , and signed this [[____]] day of [[________]] , [[______]].", space_before=6)
sig_table([
    (["**For the Institute:**", "", "Signature: ___________________________", "Name & Designation: [[____________________________]]", "Date: [[______________]]"],
     ["**Accepted by the Licensee:**", "", "Signature: ___________________________", "Name & Designation: [[____________________________]]", "(Seal of the Society)   Date: [[______________]]"]),
])

# ---- Schedule III
h1("SCHEDULE – III", center=True)
p_center("Joint Handing-Over Inventory and Condition Record", size=11.5, bold=True, space_after=2)
p_center("(To be signed on the Commencement Date and verified again at the time of return of the premises — Clauses 2.4, 14 & 19)", size=9.5, italic=True, space_after=8)
h3("A.  Institute’s fixtures / articles handed over with the premises (if any)")
table([
    ["Sl.", "Item / Fixture", "Nos.", "Condition at handing over", "Condition at taking back", "Remarks"],
    ["1", "[[____________________]]", "[[__]]", "[[________________]]", "", ""],
    ["2", "[[____________________]]", "[[__]]", "[[________________]]", "", ""],
    ["3", "[[Counters / shelves / lights / fans / switch boards — as existing]]", "[[__]]", "[[________________]]", "", ""],
    ["4", "Electrical sub-meter No. [[__________]]", "1", "Reading: [[__________]]", "Reading:", ""],
    ["5", "Water meter No. [[__________ / Not installed]]", "[[__]]", "Reading: [[__________]]", "Reading:", ""],
    ["6", "Keys (Nos.: [[__]])", "[[__]]", "[[Handed over]]", "Returned:", ""],
], widths=[0.4, 2.1, 0.6, 1.35, 1.25, 0.87], size=9.5)
h3("B.  Licensee’s equipment / articles brought into the premises")
table([
    ["Sl.", "Equipment / Article", "Make / Model / Sl. No.", "Nos.", "Date brought in", "Date removed"],
    ["1", "Coffee machine", "[[____________________________]]", "[[__]]", "[[__________]]", ""],
    ["2", "[[____________________]]", "[[____________________________]]", "[[__]]", "[[__________]]", ""],
    ["3", "[[____________________]]", "[[____________________________]]", "[[__]]", "[[__________]]", ""],
    ["4", "[[____________________]]", "[[____________________________]]", "[[__]]", "[[__________]]", ""],
], widths=[0.4, 1.9, 2.0, 0.6, 0.9, 0.77], size=9.5)
h3("C.  General condition of the premises")
p("Flooring, walls, ceiling, doors/windows, electrical and plumbing fittings found: [[in good working condition, except as noted here: ______________________________________________]].  "
  "Photographs taken on [[____________]] : [[Yes (file ref.: ________________) / No]].")
sig_table([
    (["**Handed over by (Institute):**", "", "Signature: ___________________________", "Name & Designation: [[____________________________]]", "Date: [[______________]]"],
     ["**Taken over by (Licensee):**", "", "Signature: ___________________________", "Name & Designation: [[____________________________]]", "Date: [[______________]]"]),
    (["**Taken back by (Institute) — at expiry / termination:**", "", "Signature: ___________________________", "Name & Designation: [[____________________________]]",
      "Date: [[______________]]   Meter readings: [[________________]]"],
     ["**Returned by (Licensee):**", "", "Signature: ___________________________", "Name & Designation: [[____________________________]]", "Date: [[______________]]"]),
])
p("*Remarks at taking back (repairs / recovery, if any):* [[______________________________________________________________________________]]", base_size=9.5)

# ---- Annexure A
pagebreak()
h1("ANNEXURE – A", center=True)
p_center("Extract of the Minutes of the Commercial Establishment Committee (CEC) Meeting held on 23-09-2026", size=11.5, bold=True, space_after=2)
p_center("Ref. No. 07/NITK/AC-1/2026-27/COM-ESTB/A6 dated 24-09-2026 — Item No. 4", size=10, italic=True, space_after=8)
t4 = doc.add_table(rows=1, cols=1)
t4.style = 'Table Grid'; t4.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = t4.cell(0, 0); shade(cell, LIGHT_BG); cell.paragraphs[0].clear()
paras = [
    "**Item No. 4: Proposal for Starting “Akshara Café” at E-Library**",
    "**Resolution:**",
    "(a)  The request of the NITK Consumer Co-operative Society Ltd. to run a coffee machine at “Akshara Café” in the Library premises is approved on the following terms:",
    "•  Rent: **Rs. 3,000/- (Rupees Three Thousand only) per month**",
    "•  Security Deposit: **Rs. 30,000/- (Rupees Thirty Thousand only)**",
    "•  Contract Period: **One (1) year**",
    "•  Water Charges: **Shall be charged at standard rates until a water meter is installed**",
    "(b)  The Consumer Co-operative Society shall sign an agreement with the Institute. The draft agreement copy shall be sent to the Society.",
]
for j, txt in enumerate(paras):
    para = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
    para.paragraph_format.space_after = Pt(3)
    parse_runs(para, txt, base_size=10.5)
doc.add_paragraph().paragraph_format.space_after = Pt(4)
p("*A signed copy of the complete Minutes (Ref. above) shall be enclosed with the executed Agreement. The extract above is reproduced for ready reference and forms the basis of this Licence (see Recital C).*", base_size=9.5)

# ============================================================
# PART C — FORWARDING LETTER
# ============================================================
pagebreak()
p_center("PART – C", size=12, bold=True, color=MAROON, space_after=1)
p_center("Draft Forwarding Letter to the Society (for issue after approval — Resolution 4(b))", size=11.5, bold=True, color=NAVY, space_after=6)
letterhead()

t5 = doc.add_table(rows=1, cols=2)
t5.autofit = False
c0 = t5.cell(0, 0); c1 = t5.cell(0, 1)
c0.width = Inches(4.6); c1.width = Inches(1.97)
set_cell(c0, "Ref. No.: [[07/NITK/AC-1/2026-27/COM-ESTB/____ ]]", size=10.5)
set_cell(c1, "Date: [[____-____-2026]]", size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
for row in t5.rows:
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/><w:right w:val="nil"/><w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tcBorders>')
        tcPr.append(borders)
doc.add_paragraph().paragraph_format.space_after = Pt(4)

p("To,", space_after=1)
p("**The** [[President / Secretary]],", space_after=1)
p("**NITK Consumer Co-operative Society Ltd.,**", space_after=1)
p("[[NITK Campus, Surathkal, Mangaluru – 575 025.]]", space_after=8)
p("**Subject:  Operation of coffee machine at “Akshara Café”, E-Library — Draft License Agreement — Reg.**")
p("**Reference:**  (1) Minutes of the Commercial Establishment Committee (CEC) Meeting held on 23-09-2026, Ref. No. 07/NITK/AC-1/2026-27/COM-ESTB/A6 dated 24-09-2026 (Item No. 4).  (2) Your request [[letter No. __________ dated __________]].")
p("Sir/Madam,", space_before=6)
p("With reference to the above, I am directed to state that the Commercial Establishment Committee (CEC), in its meeting held on 23-09-2026, has approved the request of the NITK Consumer Co-operative Society Ltd. to run a coffee machine at **“Akshara Café” in the E-Library premises** on the following terms: **License Fee (rent) of Rs. 3,000/- per month, Security Deposit of Rs. 30,000/-, contract period of one (1) year, and water charges at standard rates until a water meter is installed**.")
p("In terms of Resolution 4(b) of the said Minutes, the **draft License Agreement** approved by the Competent Authority, to be signed between the Institute and the Society, is enclosed herewith. You are requested to kindly peruse the same and **convey your unconditional acceptance in writing** and, within [[seven (7) / ten (10) days]] of receipt of this letter:")
numbered("(a)", "deposit the **Security Deposit of Rs. 30,000/-** [[by demand draft in favour of “The Director, NITK Surathkal”, payable at Mangaluru / NEFT as intimated]], along with the first month’s License Fee of Rs. 3,000/- [[+ taxes, if applicable]];")
numbered("(b)", "furnish copies of (i) the Society’s registration certificate, (ii) the resolution of the Society authorising the signatory of the Agreement, (iii) the valid FSSAI licence / registration for the proposed activity, (iv) PAN [[and GST registration, if any]], and (v) details of the FSSAI-registered supplier(s) for the items proposed to be sold;")
numbered("(c)", "furnish the proposed list of items with quantity and price (in the format at Schedule-II of the draft) for approval by the CEC / Competent Authority; and")
numbered("(d)", "contact the undersigned [[and the Librarian, NITK]] for joint measurement of the earmarked space (Schedule-I), completion of the remaining particulars, and **execution of the Agreement on non-judicial stamp paper of the requisite value** (purchased in the name of the Society, at its cost).")
p("Possession of the earmarked space shall be handed over **only after** execution of the Agreement, deposit of the Security Deposit and approval of the items & price list. The Society shall not install the machine or commence operations before such handing over.")
p("In case the acceptance and the documents/deposit as above are not received within the stipulated time, or the Agreement is not executed within [[fifteen (15) days]] of this letter, it will be presumed that the Society is not interested in the proposal, and the matter will be placed before the CEC for appropriate orders.")
p("Please acknowledge receipt of this letter along with the enclosed draft.")
p("", space_after=6)
para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
parse_runs(para, "Yours faithfully,", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=18)
para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
parse_runs(para, "(Signature & Seal)", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=1)
para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
parse_runs(para, "[[Assistant Registrar (Accounts) / Registrar]]", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=1)
para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
parse_runs(para, "For National Institute of Technology Karnataka, Surathkal", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10, base_size=10)
p("**Encl.:** Draft License Agreement (with Schedules I–III), as approved.")
p("**Copy to:**  1. The Director, NITK Surathkal — for kind information.  2. The Dean (FSW), Chairman, CEC — for information.  "
  "3. The Librarian, NITK Surathkal — for information and necessary action (joint measurement / handing over).  "
  "4. The Professor-in-Charge (Commercial Establishment) / (Civil Infrastructure) / (Electrical Infrastructure) — for information and necessary action (metering & billing).  "
  "5. The respective file.")

# ---- End note
line()
p_center("—  End of Draft  —", size=10, italic=True, color=GREY, space_before=6)
p("*Office note (to be deleted before execution): This DFA (Parts A–C) is submitted for approval. After approval, Part-A shall remain on file, Part-B (with all blanks filled) "
  "shall be engrossed on stamp paper for execution, and Part-C shall be issued to the Society under a formal reference number. Yellow highlights denote "
  "particulars to be filled/confirmed and must not remain in the executed Agreement.*", base_size=8.5)

doc.save(OUT)
print(f"Saved: {OUT}")
